import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Sets background shading color for a table cell."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    """Sets internal cell margins in twips."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC", sz="4"):
    """Sets clean subtle borders on a table."""
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tblBorders>
    ''')
    tblPr.append(tblBorders)

def build_agreement():
    doc = Document()

    # Set page margins
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.85)
        s.right_margin = Inches(0.85)

    # Styles
    RED_ACCENT = RGBColor(123, 8, 24)       # #7B0818
    DARK_TEXT = RGBColor(26, 26, 26)        # #1A1A1A
    MUTED_TEXT = RGBColor(80, 80, 80)       # #505050
    SLATE_BG = "F8FAFC"
    HIGHLIGHT_BG = "FFF1F2"

    # Header Box Table
    header_table = doc.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False

    # Widths: 5.0 inches and 1.8 inches
    header_table.columns[0].width = Inches(5.0)
    header_table.columns[1].width = Inches(1.8)

    cell_left = header_table.cell(0, 0)
    cell_right = header_table.cell(0, 1)

    set_cell_background(cell_left, "FFFFFF")
    set_cell_background(cell_right, "7B0818")
    set_cell_margins(cell_left, top=100, bottom=100, left=100, right=100)
    set_cell_margins(cell_right, top=140, bottom=140, left=140, right=140)

    p_title = cell_left.paragraphs[0]
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run("WEBSITE DEVELOPMENT & MAINTENANCE AGREEMENT")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(14)
    run_title.font.bold = True
    run_title.font.color.rgb = RED_ACCENT

    p_sub = cell_left.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(0)
    run_sub = p_sub.add_run("Professional 3-Year Service Agreement (36 Months) • Total Consideration: ₹2,500/-")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(9.5)
    run_sub.font.bold = True
    run_sub.font.color.rgb = MUTED_TEXT

    p_badge = cell_right.paragraphs[0]
    p_badge.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_badge.paragraph_format.space_before = Pt(4)
    p_badge.paragraph_format.space_after = Pt(0)
    run_badge = p_badge.add_run("3-YEAR CONTRACT\nAMOUNT: ₹2,500")
    run_badge.font.name = "Arial"
    run_badge.font.size = Pt(10)
    run_badge.font.bold = True
    run_badge.font.color.rgb = RGBColor(255, 255, 255)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Preamble Paragraph
    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.space_after = Pt(10)
    p_intro.paragraph_format.line_spacing = 1.2
    run_intro = p_intro.add_run(
        "This Website Development, Cloud Infrastructure & Maintenance Agreement (\"Agreement\") "
        "is made and entered into on this 17th day of September, 2026 (\"Effective Date\"), by and between the following parties:"
    )
    run_intro.font.name = "Arial"
    run_intro.font.size = Pt(10)
    run_intro.font.color.rgb = DARK_TEXT

    # Parties Table (Two-Column Cards)
    parties_table = doc.add_table(rows=1, cols=2)
    parties_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    parties_table.autofit = False
    parties_table.columns[0].width = Inches(3.35)
    parties_table.columns[1].width = Inches(3.35)
    set_table_borders(parties_table, color="D0D5DD", sz="6")

    c1 = parties_table.cell(0, 0)
    c2 = parties_table.cell(0, 1)
    set_cell_background(c1, SLATE_BG)
    set_cell_background(c2, SLATE_BG)
    set_cell_margins(c1, top=180, bottom=180, left=180, right=180)
    set_cell_margins(c2, top=180, bottom=180, left=180, right=180)

    # Party 1
    p1 = c1.paragraphs[0]
    p1.paragraph_format.space_after = Pt(6)
    r1_h = p1.add_run("1. FIRST PARTY (SERVICE PROVIDER / DEVELOPER)")
    r1_h.font.name = "Arial"
    r1_h.font.size = Pt(10)
    r1_h.font.bold = True
    r1_h.font.color.rgb = RED_ACCENT

    p1_body = c1.add_paragraph()
    p1_body.paragraph_format.space_after = Pt(2)
    p1_body.paragraph_format.line_spacing = 1.25
    r = p1_body.add_run("Name / Agency: ________________________________\n"
                       "Represented by: _______________________________\n"
                       "Designation: Full-Stack Web Engineer\n"
                       "Contact Phone: _______________________________\n"
                       "Email: ________________________________________\n"
                       "Address: ______________________________________\n"
                       "(Hereinafter referred to as the 'Service Provider')")
    r.font.name = "Arial"
    r.font.size = Pt(9.5)
    r.font.color.rgb = DARK_TEXT

    # Party 2
    p2 = c2.paragraphs[0]
    p2.paragraph_format.space_after = Pt(6)
    r2_h = p2.add_run("2. SECOND PARTY (CLIENT / WORKSHOP OWNER)")
    r2_h.font.name = "Arial"
    r2_h.font.size = Pt(10)
    r2_h.font.bold = True
    r2_h.font.color.rgb = RED_ACCENT

    p2_body = c2.add_paragraph()
    p2_body.paragraph_format.space_after = Pt(2)
    p2_body.paragraph_format.line_spacing = 1.25
    r = p2_body.add_run("Business Name: INDIAN TWO AND FOUR WHEELER\n"
                       "ALIGNMENT AND REPAIR\n"
                       "Proprietor: ____________________________________\n"
                       "Workshop: 60/1, Nehru Road, Opp. NKGSB Bank,\n"
                       "Kammanahalli, Bengaluru, Karnataka – 560084\n"
                       "Contact: 086605 20385 / +91 86605 20385\n"
                       "(Hereinafter referred to as the 'Client')")
    r.font.name = "Arial"
    r.font.size = Pt(9.5)
    r.font.color.rgb = DARK_TEXT

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Key Commercial Terms Highlight Box
    highlight_table = doc.add_table(rows=1, cols=1)
    highlight_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    highlight_table.autofit = False
    highlight_table.columns[0].width = Inches(6.8)
    set_table_borders(highlight_table, color="9A0D22", sz="8")
    hc = highlight_table.cell(0, 0)
    set_cell_background(hc, HIGHLIGHT_BG)
    set_cell_margins(hc, top=140, bottom=140, left=180, right=180)

    hp = hc.paragraphs[0]
    hp.paragraph_format.space_after = Pt(0)
    hp.paragraph_format.line_spacing = 1.25
    hr = hp.add_run(
        "• CONTRACT DURATION: 3 (Three) Consecutive Years — Effective from 17th September 2026 to 16th September 2029\n"
        "• TOTAL AGREED CONSIDERATION: ₹2,500/- (Rupees Two Thousand Five Hundred Only) All-Inclusive\n"
        "• ZERO HIDDEN CHARGES: Covers complete web hosting, Supabase cloud database, software updates, and technical support."
    )
    hr.font.name = "Arial"
    hr.font.size = Pt(9.5)
    hr.font.bold = True
    hr.font.color.rgb = RED_ACCENT

    # Helper for Section Titles
    def add_section_header(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(title)
        run.font.name = "Arial"
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = RED_ACCENT
        return p

    def add_bullet(text, bold_prefix=""):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.18
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.font.name = "Arial"
            rb.font.size = Pt(9.5)
            rb.font.bold = True
            rb.font.color.rgb = DARK_TEXT
        rt = p.add_run(text)
        rt.font.name = "Arial"
        rt.font.size = Pt(9.5)
        rt.font.color.rgb = DARK_TEXT
        return p

    # Section 1
    add_section_header("1. SCOPE OF SERVICES & SYSTEM DELIVERABLES")
    add_bullet(" High-speed Next.js 14 App Router web platform featuring dark industrial aesthetic, fully responsive across mobile smartphones, tablets, and desktop workstations.", "• Public Website & Booking Experience:")
    add_bullet(" Cloud-hosted Supabase PostgreSQL database storing customer intake logs, vehicle data, payment amounts, and automated timestamps with multi-device realtime push synchronization.", "• Cloud Database & Realtime Infrastructure:")
    add_bullet(" Dedicated 10-digit PIN-protected shop floor portal (PIN: 1234567890) for bay technicians to log cars/bikes, services done, custom jobs, and instantly generate floor receipt slips.", "• Shop Floor Worker Terminal (/worker):")
    add_bullet(" Secure 10-digit PIN-protected executive dashboard (PIN: 9876543210) calculating daily earnings, monthly accumulated revenue, yearly archives, and CSV customer export.", "• Executive Admin Analytics Dashboard (/admin):")
    add_bullet(" Daily intake counter resets automatically at 11:59:59 PM so the workshop starts each morning fresh from ₹0, while permanently archiving all records in Month and Year databases.", "• 24-Hour Midnight Rollover Engine:")
    add_bullet(" Local business rich snippets (Schema.org AutoRepair JSON-LD), dynamic sitemap.xml, robots.txt, and complete noindex protection for internal portal screens.", "• Search Engine Optimization (SEO):")

    # Section 2
    add_section_header("2. SERVICE LEVEL AGREEMENT (SLA), MAINTENANCE & SUPPORT")
    add_bullet(" The Service Provider shall ensure continuous cloud deployment with an annual platform availability target of 99.5% uptime.", "• Uptime Guarantee:")
    add_bullet(" Any critical defects, database connectivity failures, or broken intake forms will be investigated and resolved within 12 to 24 hours of notification.", "• Technical Defect Resolution:")
    add_bullet(" The Client is entitled to request routine text modifications, service price updates, or photo swaps up to four (4) times per calendar year at no additional charge.", "• Quarterly Revisions:")
    add_bullet(" Automated daily checks guarantee zero demo orders pollute the live business account and genuine customer data is permanently retained.", "• Database Health & Zero-Demo Protection:")

    # Section 3
    add_section_header("3. DATA OWNERSHIP, CONFIDENTIALITY & INTELLECTUAL PROPERTY")
    add_bullet(" The Client retains 100% full and exclusive ownership of all customer records, vehicle registration numbers, job histories, financial amounts, and brand identity.", "• Exclusive Client Ownership:")
    add_bullet(" The Service Provider shall maintain strict confidentiality regarding portal access passwords (Admin/Worker PINs), daily workshop earnings, and customer phone numbers.", "• Data Privacy:")
    add_bullet(" The Service Provider is strictly prohibited from selling, sharing, or exposing Client business data to any third party.", "• Non-Disclosure:")

    # Section 4
    add_section_header("4. COMMERCIAL TERMS & PAYMENT CONFIRMATION")
    p_comm = doc.add_paragraph()
    p_comm.paragraph_format.space_before = Pt(2)
    p_comm.paragraph_format.space_after = Pt(4)
    p_comm.paragraph_format.line_spacing = 1.2
    r_comm = p_comm.add_run(
        "The total agreed contract value for the full 36-month term is ₹2,500/- (Rupees Two Thousand Five Hundred Only).\n"
        "Payment Mode: [ ] Cash    [ ] UPI / GPay / PhonePe    [ ] Bank Transfer (NEFT/IMPS)\n"
        "Payment Reference / UTR No.: ___________________________________    Date: ________________________\n"
        "Payment Status: [ ] Received in Full    [ ] Part Paid (Balance: ₹___________)"
    )
    r_comm.font.name = "Arial"
    r_comm.font.size = Pt(9.5)
    r_comm.font.color.rgb = DARK_TEXT

    # Section 5: Signatures
    add_section_header("5. SIGNATURES & EXECUTION")
    p_sign_intro = doc.add_paragraph()
    p_sign_intro.paragraph_format.space_after = Pt(8)
    r_si = p_sign_intro.add_run("IN WITNESS WHEREOF, the Parties hereto have executed this Agreement as of the Effective Date written above:")
    r_si.font.name = "Arial"
    r_si.font.size = Pt(9.5)

    sig_table = doc.add_table(rows=1, cols=2)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_table.autofit = False
    sig_table.columns[0].width = Inches(3.35)
    sig_table.columns[1].width = Inches(3.35)
    set_table_borders(sig_table, color="B0B8C5", sz="6")

    sc1 = sig_table.cell(0, 0)
    sc2 = sig_table.cell(0, 1)
    set_cell_margins(sc1, top=140, bottom=140, left=140, right=140)
    set_cell_margins(sc2, top=140, bottom=140, left=140, right=140)

    # Sig Provider
    sp1 = sc1.paragraphs[0]
    sp1.paragraph_format.space_after = Pt(4)
    r = sp1.add_run("FOR SERVICE PROVIDER / DEVELOPER:")
    r.font.name = "Arial"
    r.font.size = Pt(9.5)
    r.font.bold = True
    r.font.color.rgb = RED_ACCENT

    sp1_b = sc1.add_paragraph()
    sp1_b.paragraph_format.space_after = Pt(0)
    sp1_b.paragraph_format.line_spacing = 1.3
    r = sp1_b.add_run(
        "Name: ____________________________________\n"
        "Designation: Full-Stack Web Engineer\n\n\n"
        "Authorized Signature: ______________________\n"
        "Date: 17th September 2026\n"
        "Place: Bengaluru, Karnataka"
    )
    r.font.name = "Arial"
    r.font.size = Pt(9)

    # Sig Client
    sp2 = sc2.paragraphs[0]
    sp2.paragraph_format.space_after = Pt(4)
    r = sp2.add_run("FOR CLIENT / WORKSHOP OWNER:")
    r.font.name = "Arial"
    r.font.size = Pt(9.5)
    r.font.bold = True
    r.font.color.rgb = RED_ACCENT

    sp2_b = sc2.add_paragraph()
    sp2_b.paragraph_format.space_after = Pt(0)
    sp2_b.paragraph_format.line_spacing = 1.3
    r = sp2_b.add_run(
        "Business: Indian Two & Four Wheeler Alignment\n"
        "Proprietor: _______________________________\n\n\n"
        "Signature & Stamp: ________________________\n"
        "Date: 17th September 2026\n"
        "Place: Kammanahalli, Bengaluru"
    )
    r.font.name = "Arial"
    r.font.size = Pt(9)

    # Witnesses
    p_wit = doc.add_paragraph()
    p_wit.paragraph_format.space_before = Pt(10)
    p_wit.paragraph_format.space_after = Pt(4)
    rw = p_wit.add_run("Witnesses:")
    rw.font.name = "Arial"
    rw.font.size = Pt(9.5)
    rw.font.bold = True

    wit_table = doc.add_table(rows=1, cols=2)
    wit_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    wit_table.autofit = False
    wit_table.columns[0].width = Inches(3.35)
    wit_table.columns[1].width = Inches(3.35)
    set_table_borders(wit_table, color="D0D5DD", sz="4")

    wc1 = wit_table.cell(0, 0)
    wc2 = wit_table.cell(0, 1)
    set_cell_margins(wc1, top=100, bottom=100, left=120, right=120)
    set_cell_margins(wc2, top=100, bottom=100, left=120, right=120)

    wp1 = wc1.paragraphs[0]
    wp1.paragraph_format.line_spacing = 1.25
    r = wp1.add_run("1. Witness Name: _________________________\n"
                    "   Address: ______________________________\n"
                    "   Signature: ____________________________")
    r.font.name = "Arial"
    r.font.size = Pt(8.5)

    wp2 = wc2.paragraphs[0]
    wp2.paragraph_format.line_spacing = 1.25
    r = wp2.add_run("2. Witness Name: _________________________\n"
                    "   Address: ______________________________\n"
                    "   Signature: ____________________________")
    r.font.name = "Arial"
    r.font.size = Pt(8.5)

    output_path = os.path.abspath("WEBSITE_SERVICE_AGREEMENT.docx")
    doc.save(output_path)
    print(f"Successfully generated DOCX at: {output_path}")

if __name__ == "__main__":
    build_agreement()
